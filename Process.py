"""
Monthly Trade Volume Summary Generator
========================================
Reads all order-audit Excel files (the "partXofY" files) from the "Excel Files"
folder, aggregates buy/sell trade volume (Qty x Gross Fill Price) by ISIN for
each calendar month, and writes one output workbook to the "Output" folder.

- Grouping key: ISIN (column D). Instrument Description (column C) is carried
  along for readability. Differences in Instrument Code (column B) - e.g.
  buy-in / cash market / main market variants - are ignored, per spec.
- Month is derived from Trade Date (column AB, format yyyy-mm-dd).
- September is always excluded from the summary.
- Output: one sheet ("Monthly Volume Summary") with one stacked table per
  month (ISIN | Instrument Description | Buy Volume (IDR) | Sell Volume (IDR)),
  plus a second sheet ("USD-IDR FX Reference") listing the average USD/IDR
  rate for each month shown, for reference only (not used in the calculation
  since all trade values are already in IDR).

HOW TO RUN
----------
1. Place this script in a folder that also contains an "Excel Files"
   subfolder with the 8 (or however many) part files, and an "Output"
   subfolder (created automatically if missing).
2. Run:  python generate_summary.py
3. The result appears in Output/Monthly_Trade_Volume_Summary_<timestamp>.xlsx

CONFIG
------
Adjust INPUT_FOLDER / OUTPUT_FOLDER below if your folders live elsewhere.
Adjust FX_RATES_BY_MONTH if your data spans months not already listed there.
"""

import glob
import os
import sys
import datetime as dt
from collections import defaultdict

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# ------------------------------------------------------------------------
# CONFIG
# ------------------------------------------------------------------------
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
INPUT_FOLDER = os.path.join(SCRIPT_DIR, "Excel Files")
OUTPUT_FOLDER = os.path.join(SCRIPT_DIR, "Output")

# Fixed column positions (0-indexed) matching the audit file layout.
# Column letters as given in the spec: B, C, D, N, P, R, AB
COL_INSTRUMENT_CODE = 1    # B
COL_INSTRUMENT_DESC = 2    # C
COL_ISIN = 3               # D
COL_SIDE = 13              # N
COL_QTY_FILLED = 15        # P
COL_GROSS_PRICE = 17       # R
COL_TRADE_DATE = 27        # AB

EXCLUDED_MONTH = 9  # September is always dropped from the summary

# Average USD -> IDR rate per month, pulled from x-rates.com monthly averages
# (interbank/market mid rate) on 2026-09-14. Update/add entries here if your
# data covers months not listed - the script will flag any month it can't
# find a rate for instead of guessing.
FX_RATES_BY_MONTH = {
    "2026-01": 16807.31,
    "2026-02": 16828.17,
    "2026-03": 16919.46,
    "2026-04": 17123.89,
    "2026-05": 17563.99,
    "2026-06": 17910.47,
    "2026-07": 18019.54,
    "2026-08": 17831.21,
}
FX_SOURCE_NOTE = "Source: x-rates.com monthly average, USD/IDR (retrieved 2026-09-14)"

# ------------------------------------------------------------------------
# STEP 1: Load and combine all input files
# ------------------------------------------------------------------------

def find_input_files(folder):
    if not os.path.isdir(folder):
        sys.exit(f"ERROR: Input folder not found: {folder}")
    files = [
        f for f in glob.glob(os.path.join(folder, "*.xlsx"))
        if not os.path.basename(f).startswith("~$")
    ]
    if not files:
        sys.exit(f"ERROR: No .xlsx files found in: {folder}")
    return sorted(files)


def load_file(path):
    """Read one audit file and return a cleaned DataFrame with just the
    columns we need, using fixed column positions (not header names, since
    header text may vary slightly between files)."""
    raw = pd.read_excel(path, header=0)

    needed_max_idx = max(COL_INSTRUMENT_CODE, COL_INSTRUMENT_DESC, COL_ISIN,
                          COL_SIDE, COL_QTY_FILLED, COL_GROSS_PRICE, COL_TRADE_DATE)
    if raw.shape[1] <= needed_max_idx:
        print(f"  WARNING: '{os.path.basename(path)}' has only {raw.shape[1]} "
              f"columns (need at least {needed_max_idx + 1}). Skipping file.")
        return None

    df = pd.DataFrame({
        "InstrumentCode": raw.iloc[:, COL_INSTRUMENT_CODE],
        "InstrumentDescription": raw.iloc[:, COL_INSTRUMENT_DESC],
        "ISIN": raw.iloc[:, COL_ISIN],
        "Side": raw.iloc[:, COL_SIDE],
        "Qty": pd.to_numeric(raw.iloc[:, COL_QTY_FILLED], errors="coerce"),
        "Price": pd.to_numeric(raw.iloc[:, COL_GROSS_PRICE], errors="coerce"),
        "TradeDate": pd.to_datetime(raw.iloc[:, COL_TRADE_DATE], errors="coerce"),
    })

    before = len(df)
    df = df.dropna(subset=["ISIN", "Side", "Qty", "Price", "TradeDate"])
    dropped = before - len(df)
    if dropped:
        print(f"  Note: '{os.path.basename(path)}' - dropped {dropped} row(s) "
              f"with missing/invalid ISIN, Side, Qty, Price, or Trade Date.")

    df["Side"] = df["Side"].astype(str).str.strip().str.upper()
    df = df[df["Side"].isin(["B", "S"])]
    df["Volume"] = df["Qty"] * df["Price"]
    df["Month"] = df["TradeDate"].dt.strftime("%Y-%m")
    df["SourceFile"] = os.path.basename(path)
    return df


def load_all_files(folder):
    files = find_input_files(folder)
    print(f"Found {len(files)} input file(s) in '{folder}':")
    for f in files:
        print(f"  - {os.path.basename(f)}")

    frames = []
    for f in files:
        df = load_file(f)
        if df is not None and len(df):
            frames.append(df)
    if not frames:
        sys.exit("ERROR: No usable rows found across all input files.")
    combined = pd.concat(frames, ignore_index=True)
    return combined


# ------------------------------------------------------------------------
# STEP 2: Aggregate by month / ISIN / side
# ------------------------------------------------------------------------

def build_monthly_tables(df):
    """Returns dict: month ('YYYY-MM') -> DataFrame[ISIN, Description, Buy, Sell]
    sorted by month, with September excluded."""
    df = df[df["TradeDate"].dt.month != EXCLUDED_MONTH].copy()

    grouped = (
        df.groupby(["Month", "ISIN", "InstrumentDescription", "Side"])["Volume"]
        .sum()
        .reset_index()
    )

    tables = {}
    for month, month_df in grouped.groupby("Month"):
        pivot = month_df.pivot_table(
            index=["ISIN", "InstrumentDescription"],
            columns="Side",
            values="Volume",
            aggfunc="sum",
            fill_value=0,
        ).reset_index()

        if "B" not in pivot.columns:
            pivot["B"] = 0
        if "S" not in pivot.columns:
            pivot["S"] = 0

        pivot = pivot.rename(columns={
            "InstrumentDescription": "Instrument Description",
            "B": "Buy Volume (IDR)",
            "S": "Sell Volume (IDR)",
        })
        pivot = pivot[["ISIN", "Instrument Description", "Buy Volume (IDR)", "Sell Volume (IDR)"]]
        pivot = pivot.sort_values("ISIN").reset_index(drop=True)
        tables[month] = pivot

    return dict(sorted(tables.items()))


# ------------------------------------------------------------------------
# STEP 3: Write output workbook
# ------------------------------------------------------------------------

MONTH_NAMES = {
    "01": "January", "02": "February", "03": "March", "04": "April",
    "05": "May", "06": "June", "07": "July", "08": "August",
    "09": "September", "10": "October", "11": "November", "12": "December",
}


def month_label(month_key):
    year, mm = month_key.split("-")
    return f"{MONTH_NAMES.get(mm, mm)} {year}"


def write_output(tables, out_path):
    wb = Workbook()

    # --- Sheet 1: Monthly Volume Summary ---
    ws = wb.active
    ws.title = "Monthly Volume Summary"

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    month_font = Font(name="Arial", bold=True, size=13)
    normal_font = Font(name="Arial")
    col_widths = [18, 38, 22, 22]

    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    row = 1
    for month_key, table in tables.items():
        ws.cell(row=row, column=1, value=month_label(month_key)).font = month_font
        row += 1

        headers = ["ISIN", "Instrument Description", "Buy Volume (IDR)", "Sell Volume (IDR)"]
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=c, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        row += 1

        for _, r in table.iterrows():
            ws.cell(row=row, column=1, value=r["ISIN"]).font = normal_font
            ws.cell(row=row, column=2, value=r["Instrument Description"]).font = normal_font
            buy_cell = ws.cell(row=row, column=3, value=float(r["Buy Volume (IDR)"]))
            sell_cell = ws.cell(row=row, column=4, value=float(r["Sell Volume (IDR)"]))
            buy_cell.font = normal_font
            sell_cell.font = normal_font
            buy_cell.number_format = "#,##0"
            sell_cell.number_format = "#,##0"
            row += 1

        # Totals row for the month
        total_row = row
        ws.cell(row=total_row, column=2, value="Total").font = Font(name="Arial", bold=True)
        first_data_row = total_row - len(table)
        last_data_row = total_row - 1
        if len(table):
            ws.cell(row=total_row, column=3,
                    value=f"=SUM(C{first_data_row}:C{last_data_row})").font = Font(name="Arial", bold=True)
            ws.cell(row=total_row, column=4,
                    value=f"=SUM(D{first_data_row}:D{last_data_row})").font = Font(name="Arial", bold=True)
        ws.cell(row=total_row, column=3).number_format = "#,##0"
        ws.cell(row=total_row, column=4).number_format = "#,##0"
        row += 2  # blank row separator

    # --- Sheet 2: USD-IDR FX Reference ---
    fx_ws = wb.create_sheet("USD-IDR FX Reference")
    fx_ws.column_dimensions["A"].width = 14
    fx_ws.column_dimensions["B"].width = 22
    fx_ws.column_dimensions["C"].width = 60

    fx_ws.cell(row=1, column=1, value="Month").font = header_font
    fx_ws.cell(row=1, column=2, value="Avg 1 USD = IDR").font = header_font
    fx_ws.cell(row=1, column=3, value="Notes").font = header_font
    for c in (1, 2, 3):
        fx_ws.cell(row=1, column=c).fill = header_fill

    r = 2
    for month_key in tables.keys():
        fx_ws.cell(row=r, column=1, value=month_label(month_key)).font = normal_font
        rate = FX_RATES_BY_MONTH.get(month_key)
        if rate is None:
            fx_ws.cell(row=r, column=2, value="N/A - add rate manually").font = normal_font
        else:
            cell = fx_ws.cell(row=r, column=2, value=rate)
            cell.font = normal_font
            cell.number_format = "#,##0.00"
        fx_ws.cell(row=r, column=3, value=FX_SOURCE_NOTE).font = Font(name="Arial", italic=True, size=9)
        r += 1

    wb.save(out_path)


# ------------------------------------------------------------------------
# MAIN
# ------------------------------------------------------------------------

def main():
    os.makedirs(OUTPUT_FOLDER, exist_ok=True)

    print(f"Reading audit files from: {INPUT_FOLDER}")
    combined = load_all_files(INPUT_FOLDER)
    print(f"Total rows loaded: {len(combined)}")

    tables = build_monthly_tables(combined)
    if not tables:
        sys.exit("ERROR: No data left after excluding September / invalid rows.")

    print(f"Months included: {', '.join(tables.keys())} (September excluded)")

    missing_fx = [m for m in tables if m not in FX_RATES_BY_MONTH]
    if missing_fx:
        print(f"  NOTE: No FX rate on file for: {', '.join(missing_fx)}. "
              f"These will show 'N/A' on the FX Reference sheet - add manually if needed.")

    timestamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = os.path.join(OUTPUT_FOLDER, f"Monthly_Trade_Volume_Summary_{timestamp}.xlsx")
    write_output(tables, out_path)

    print(f"\nDone. Output written to:\n  {out_path}")


if __name__ == "__main__":
    main()
